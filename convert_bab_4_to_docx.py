import re
import os
from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner margins (padding) for a table cell in dxa (1 pt = 20 dxa)"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    """Set background color (shading) for a cell"""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def add_inline_formatting(paragraph, text, is_code=False, font_size=12):
    """Parse simple markdown bold (**), italic (*), and inline code (`) and add runs to a paragraph"""
    if is_code:
        run = paragraph.add_run(text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        return
        
    # Pattern to find markdown markers
    # We will search for **, *, and `
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
        else:
            run = paragraph.add_run(part)
        
        # Ensure base font is Times New Roman if not consolas
        if run.font.name != 'Consolas':
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)

def convert_md_to_docx(md_path, docx_path):
    print(f"Reading Markdown from: {md_path}")
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    doc = Document()
    
    # ---------------------------------------------------------
    # PENGATURAN HALAMAN SKRIPSI STANDAR INDONESIA
    # Kertas: A4
    # Margin: Kiri = 4 cm, Atas = 3 cm, Kanan = 3 cm, Bawah = 3 cm
    # ---------------------------------------------------------
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.0)
    section.bottom_margin = Cm(3.0)
    section.left_margin = Cm(4.0)
    section.right_margin = Cm(3.0)
    
    # Set default style to Times New Roman, 12pt, 1.5 line spacing, Justified
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    p_format = style.paragraph_format
    p_format.line_spacing = 1.5
    p_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_format.space_after = Pt(6)
    
    in_code_block = False
    code_content = []
    
    in_table = False
    table_rows = []
    
    for line in lines:
        line_stripped = line.strip()
        
        # ---------------------------------------------------------
        # PENGATURAN BLOK KODE (Fenced Code Block ```)
        # ---------------------------------------------------------
        if line_stripped.startswith('```'):
            if in_code_block:
                # End of code block, render it
                in_code_block = False
                code_text = "".join(code_content)
                
                # Add table container for code block to draw light borders and gray background
                tbl = doc.add_table(rows=1, cols=1)
                tbl.autofit = False
                cell = tbl.cell(0, 0)
                cell.width = Cm(14.0) # fit page width (21cm - 4cm left - 3cm right = 14cm)
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                set_cell_shading(cell, "F1F5F9") # Slate 100 hex color
                
                cp = cell.paragraphs[0]
                cp.paragraph_format.line_spacing = 1.0
                cp.paragraph_format.space_after = Pt(0)
                
                add_inline_formatting(cp, code_text, is_code=True)
                
                # Spacer paragraph after table
                spacer = doc.add_paragraph()
                spacer.paragraph_format.line_spacing = 1.0
                spacer.paragraph_format.space_after = Pt(6)
                spacer.paragraph_format.space_before = Pt(0)
                spacer.paragraph_format.height = Pt(4)
                
                code_content = []
            else:
                in_code_block = True
            continue
            
        if in_code_block:
            code_content.append(line)
            continue
            
        # ---------------------------------------------------------
        # PENGATURAN TABEL (Markdown Table)
        # ---------------------------------------------------------
        if line_stripped.startswith('|'):
            # It's a table row
            # Ignore separator rows like |---|---|
            if re.match(r'^\|[\s\-\|]+$', line_stripped):
                continue
            
            # Parse row cells
            cells_data = [c.strip() for c in line_stripped.split('|')[1:-1]]
            table_rows.append(cells_data)
            in_table = True
            continue
        else:
            if in_table:
                # End of table, render accumulated table rows
                in_table = False
                if table_rows:
                    num_cols = len(table_rows[0])
                    num_rows = len(table_rows)
                    tbl = doc.add_table(rows=num_rows, cols=num_cols)
                    tbl.style = 'Table Grid'
                    
                    # Style table
                    for r_idx, row_data in enumerate(table_rows):
                        row = tbl.rows[r_idx]
                        
                        # Set height or padding
                        for c_idx, cell_value in enumerate(row_data):
                            if c_idx >= num_cols:
                                break
                            cell = row.cells[c_idx]
                            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                            
                            # Header shading
                            if r_idx == 0:
                                set_cell_shading(cell, "4B176E") # Tema ungu EduConnect
                                
                            cp = cell.paragraphs[0]
                            cp.paragraph_format.line_spacing = 1.15
                            cp.paragraph_format.space_after = Pt(2)
                            cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                            
                            # Add text with formatting
                            if r_idx == 0:
                                run = cp.add_run(cell_value)
                                run.bold = True
                                run.font.name = 'Times New Roman'
                                run.font.size = Pt(10)
                                run.font.color.rgb = docx.shared.RGBColor(255, 255, 255) # Putih untuk header
                            else:
                                add_inline_formatting(cp, cell_value, font_size=10)
                                
                    # Add space after table
                    doc.add_paragraph().paragraph_format.space_after = Pt(8)
                table_rows = []
                
        if not line_stripped:
            continue
            
        # ---------------------------------------------------------
        # PENGATURAN HEADINGS
        # ---------------------------------------------------------
        # H1 (BAB Title) - Contoh: # BAB IV
        if line_stripped.startswith('# '):
            heading_text = line_stripped[2:]
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.line_spacing = 1.5
            
            run = p.add_run(heading_text.upper())
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            continue
            
        # H2 (Sub-bab Utama) - Contoh: ## C. Hasil Implementasi ...
        if line_stripped.startswith('## '):
            heading_text = line_stripped[3:]
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.keep_with_next = True
            
            run = p.add_run(heading_text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            continue
            
        # H3 (Sub-sub-bab) - Contoh: ### 1. Hasil Implementasi Sistem
        if line_stripped.startswith('### '):
            heading_text = line_stripped[4:]
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            
            run = p.add_run(heading_text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            continue
            
        # H4 (Sub-sub-sub-bab / List Bold) - Contoh: #### a. Halaman Login dan Registrasi
        if line_stripped.startswith('#### '):
            heading_text = line_stripped[5:]
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            
            run = p.add_run(heading_text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            continue
            
        # ---------------------------------------------------------
        # PENGATURAN LISTS
        # ---------------------------------------------------------
        # Bullet list - Contoh: * item
        if line_stripped.startswith('* '):
            list_text = line_stripped[2:]
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.5)
            
            add_inline_formatting(p, list_text)
            continue
            
        # Numbered list - Contoh: 1. item
        numbered_match = re.match(r'^(\d+)\.\s(.*)$', line_stripped)
        if numbered_match:
            num = numbered_match.group(1)
            list_text = numbered_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.5)
            
            # Simple numbered run
            run_num = p.add_run(f"{num}. ")
            run_num.font.name = 'Times New Roman'
            run_num.font.size = Pt(12)
            
            add_inline_formatting(p, list_text)
            continue
            
        # ---------------------------------------------------------
        # PARAGRAF BIASA
        # ---------------------------------------------------------
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline_formatting(p, line_stripped)
        
    print(f"Saving DOCX to: {docx_path}")
    doc.save(docx_path)
    print("DOCX successfully saved!")

if __name__ == "__main__":
    md_file = r"C:\Users\USER\.gemini\antigravity\brain\271748ec-f344-4f7d-bbae-75aa9ba0718a\bab_4_hasil_dan_pembahasan.md"
    docx_file = r"d:\apk_edu\BAB_IV_Hasil_dan_Pembahasan_EduConnect.docx"
    convert_md_to_docx(md_file, docx_file)
