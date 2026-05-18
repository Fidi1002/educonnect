from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_skripsi_document():
    doc = Document()
    
    # Judul Dokumen
    title = doc.add_heading('DOKUMENTASI FINAL SKRIPSI - EDUCONNECT', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph('Dokumen ini berisi rangkuman fungsionalitas, pengujian, dan arsitektur sistem dari proyek EduConnect yang dibangun dengan Flutter dan Supabase.\n')

    # ==========================================
    # 1. DOKUMENTASI FITUR
    # ==========================================
    doc.add_heading('1. Dokumentasi Fitur dan Kebutuhan Screenshot', level=1)
    doc.add_paragraph('Berikut adalah daftar fitur utama beserta tangkapan layar (screenshot) yang perlu dilampirkan pada laporan skripsi:')
    
    table_fitur = doc.add_table(rows=1, cols=5)
    table_fitur.style = 'Table Grid'
    hdr_cells = table_fitur.rows[0].cells
    hdr_cells[0].text = 'No'
    hdr_cells[1].text = 'Nama Fitur'
    hdr_cells[2].text = 'Deskripsi'
    hdr_cells[3].text = 'File/Modul Terkait'
    hdr_cells[4].text = 'Kebutuhan Screenshot'

    fitur_data = [
        ('1', 'Autentikasi & Role', 'Registrasi/login dan pemilihan peran (Student/Tutor).', 'lib/features/auth/', 'Halaman Login, Form Register, Pemilihan Role'),
        ('2', 'Profil & Ketersediaan', 'Tutor mengisi profil, upload foto, dan mengatur ketersediaan waktu.', 'lib/features/tutor_profile/', 'Form Profil Tutor, Form Kalender Ketersediaan'),
        ('3', 'Pencarian Radius', 'Mencari tutor terdekat dalam radius 1/5/10/20 km berbasis PostGIS.', 'lib/features/home/', 'Peta/List Tutor, Slider Radius Lokasi'),
        ('4', 'Booking & Pembayaran', 'Pemesanan paket belajar (1-6 bulan) dan simulasi pembayaran (dummy).', 'lib/features/booking/', 'Detail Tutor, Form Booking, Transaksi Sukses'),
        ('5', 'Sesi Belajar & PR', 'Manajemen sesi, reminder otomatis, dan pengumpulan pekerjaan rumah (PR).', 'lib/features/home/', 'Kalender Sesi, Tombol Konfirmasi, Form Submit PR'),
        ('6', 'Real-time Chat', 'Komunikasi instan antar pengguna berdasarkan booking aktif.', 'lib/features/chat/', 'Daftar Obrolan, Ruang Chat (Balon Pesan)')
    ]

    for no, nama, desk, file_mod, ss in fitur_data:
        row_cells = table_fitur.add_row().cells
        row_cells[0].text = no
        row_cells[1].text = nama
        row_cells[2].text = desk
        row_cells[3].text = file_mod
        row_cells[4].text = ss

    doc.add_paragraph('\n')

    # ==========================================
    # 2. SKENARIO PENGUJIAN FINAL
    # ==========================================
    doc.add_heading('2. Skenario Pengujian Final (Black Box)', level=1)
    doc.add_paragraph('Pengujian perangkat lunak dilakukan menggunakan metode Black Box Testing yang berfokus pada alur input-output sistem:')
    
    table_test = doc.add_table(rows=1, cols=6)
    table_test.style = 'Table Grid'
    hdr_test = table_test.rows[0].cells
    headers = ['No', 'Skenario Uji', 'Langkah Pengujian', 'Hasil Diharapkan', 'Hasil Aktual', 'Status']
    for i, header in enumerate(headers):
        hdr_test[i].text = header

    test_data = [
        ('1', 'Registrasi akun murid', 'Isi nama, email, password -> Register', 'Akun terbuat & diarahkan ke onboarding', 'Sesuai', 'Lulus'),
        ('2', 'Pencarian tutor radius', 'Atur radius ke 5 km di Beranda', 'Tutor dalam radius 5 km tampil di list', 'Sesuai', 'Lulus'),
        ('3', 'Booking kuota penuh (Negatif)', 'Booking tutor dengan > 2 murid aktif', 'Ditolak oleh Trigger Supabase (Error UI)', 'Sesuai', 'Lulus'),
        ('4', 'Pembayaran Simulasi', 'Klik "Bayar (Dummy)" pada pesanan', 'Status pesanan berubah menjadi Paid', 'Sesuai', 'Lulus'),
        ('5', 'Kirim pesan realtime', 'Ketik teks di chat room -> Kirim', 'Pesan terkirim ke lawan bicara tanpa refresh', 'Sesuai', 'Lulus'),
        ('6', 'Notifikasi H-24 (Push)', 'Jalankan edge function process-reminders', 'Notifikasi push masuk ke perangkat', 'Sesuai', 'Lulus'),
        ('7', 'Transisi ilegal (Negatif)', 'Murid bypass API selesaikan transaksi', 'Akses ditolak oleh RLS Database', 'Sesuai', 'Lulus')
    ]

    for no, skenario, langkah, ekspektasi, aktual, status in test_data:
        row_cells = table_test.add_row().cells
        row_cells[0].text = no
        row_cells[1].text = skenario
        row_cells[2].text = langkah
        row_cells[3].text = ekspektasi
        row_cells[4].text = aktual
        row_cells[5].text = status

    doc.add_paragraph('\n')

    # ==========================================
    # 3. DEMO ACCOUNT
    # ==========================================
    doc.add_heading('3. Akun Pengujian (Demo Account)', level=1)
    doc.add_paragraph('Berikut adalah konfigurasi akun demo yang telah disiapkan di Supabase Auth untuk memudahkan proses pengujian saat sidang:')
    
    table_demo = doc.add_table(rows=1, cols=4)
    table_demo.style = 'Table Grid'
    hdr_demo = table_demo.rows[0].cells
    hdr_demo[0].text = 'Role'
    hdr_demo[1].text = 'Email'
    hdr_demo[2].text = 'Password'
    hdr_demo[3].text = 'Keterangan Uji'

    demo_data = [
        ('Student', 'murid1@educonnect.com', 'password123', 'Memiliki akses pencarian lokasi & saldo bayar'),
        ('Student', 'murid2@educonnect.com', 'password123', 'Digunakan untuk tes bentrok booking'),
        ('Tutor', 'tutor1@educonnect.com', 'password123', 'Profil lengkap, siap terima booking & chat')
    ]

    for role, email, password, ket in demo_data:
        row_cells = table_demo.add_row().cells
        row_cells[0].text = role
        row_cells[1].text = email
        row_cells[2].text = password
        row_cells[3].text = ket

    doc.add_paragraph('\n')

    # ==========================================
    # 4. NARASI ARSITEKTUR & PENGUJIAN
    # ==========================================
    doc.add_heading('4. Narasi Arsitektur Sistem', level=1)
    doc.add_paragraph('Aplikasi EduConnect dibangun menggunakan arsitektur modular berbasis komponen modern (Feature-First Modular Architecture) pada sisi frontend dan pendekatan Backend-as-a-Service (BaaS) menggunakan platform Supabase. Pada sisi Frontend, aplikasi dikembangkan menggunakan bahasa pemrograman Dart dengan kerangka kerja (framework) Flutter. Kode program dipecah berdasarkan fitur, di mana masing-masing fitur memiliki layer Domain, Data, Application, dan Presentation. Pola Dependency Injection diterapkan melalui Riverpod Providers.')
    doc.add_paragraph('Pada sisi Backend, sistem menggunakan PostgreSQL. Arsitektur data menerapkan pembatasan hak akses Row Level Security (RLS). Seluruh aturan bisnis krusial, seperti pembatasan kuota pesanan, pencegahan bentrok jadwal, dan transisi status, divalidasi langsung di layer database melalui Database Triggers dan Remote Procedure Calls (RPC) untuk menjamin persistensi dan keamanan data dari intervensi API yang tidak sah.')

    doc.add_heading('5. Narasi Pengujian Sistem', level=1)
    doc.add_paragraph('Tahap pengujian dilakukan menggunakan metode Black Box Testing dipadukan dengan End-to-End (E2E) Regression Suite. Pengujian Black Box menitikberatkan pada evaluasi fungsionalitas sistem berdasarkan masukan (input) dan keluaran (output). Pengujian diklasifikasikan ke dalam uji coba positif untuk memverifikasi proses sukses, dan uji coba negatif untuk mengukur kemampuan sistem menahan percobaan intervensi ilegal (Illegitimate State Transition). Hasil pengujian mencatatkan bahwa sistem telah mencapai tingkat keberhasilan 100% pada skenario yang dirancang, memastikan keandalan manajemen status pada level RDBMS.')

    # ==========================================
    # 6. TEKNOLOGI YANG DIGUNAKAN
    # ==========================================
    doc.add_heading('6. Analisis Teknologi', level=1)
    table_tech = doc.add_table(rows=1, cols=2)
    table_tech.style = 'Table Grid'
    hdr_tech = table_tech.rows[0].cells
    hdr_tech[0].text = 'Teknologi'
    hdr_tech[1].text = 'Peran / Fungsi'

    tech_data = [
        ('Flutter & Dart', 'Pengembangan antarmuka lintas platform (Android/Web).'),
        ('Supabase', 'Platform BaaS yang menangani Database, Auth, dan Storage.'),
        ('PostgreSQL & PostGIS', 'Basis data relasional dengan ekstensi spasial untuk kalkulasi radius.'),
        ('Riverpod & GoRouter', 'State management reaktif dan sistem pengamanan routing (Auth Guard).'),
        ('Deno (Edge Functions)', 'Runtime serverless untuk mengeksekusi cron jobs reminder sesi.'),
        ('Firebase Cloud Messaging', 'Infrastruktur pengiriman push notification ke device (HTTP v1 API).')
    ]

    for tech, fungsi in tech_data:
        row_cells = table_tech.add_row().cells
        row_cells[0].text = tech
        row_cells[1].text = fungsi

    doc.add_paragraph('\n')

    # ==========================================
    # 7. FLOW APLIKASI
    # ==========================================
    doc.add_heading('7. Flow Aplikasi (Alur Penggunaan Utama)', level=1)
    doc.add_paragraph('1. Autentikasi: Buka Aplikasi -> Input Kredensial -> Pilih Role (Student/Tutor) -> Beranda.')
    doc.add_paragraph('2. Transaksi (Murid): Izinkan Lokasi -> Atur Radius -> Pilih Tutor -> Pilih Paket (1-6 bulan) & Jadwal -> Klik Booking.')
    doc.add_paragraph('3. Validasi (Tutor): Terima Notifikasi Booking Masuk -> Buka Detail -> Terima (Accept).')
    doc.add_paragraph('4. Pembayaran & Eksekusi: Murid Bayar (Dummy) -> Transaksi "Paid" -> Sesi Belajar Terbentuk -> Pelaksanaan Sesi -> Tutor Tandai Selesai -> Murid Konfirmasi Hadir -> Isi PR/Materi.')

    # Simpan File
    filename = 'Dokumentasi_Skripsi_EduConnect.docx'
    doc.save(filename)
    print(f'Dokumen berhasil dibuat dan disimpan dengan nama: {filename}')

if __name__ == '__main__':
    create_skripsi_document()